#!/usr/bin/env python3
"""
excel_to_seed.py — one-off: convierte los Bloques A/B/C de la hoja "⚙️ Parámetros"
del Excel de Samavi en un seed SQL para Supabase (tablas listings / general_expenses / events).

Uso:
    python3 scripts/excel_to_seed.py ["STAG SAMAVI — Dashboard 2026.xlsx"] [supabase/seed/seed.sql]

No toca la nube: solo lee el Excel y escribe un .sql idempotente (TRUNCATE + INSERT).
"""
import sys
import datetime as dt
import openpyxl

DEFAULT_XLSX = "STAG SAMAVI — Dashboard 2026.xlsx"
DEFAULT_OUT = "supabase/seed/seed.sql"
PARAMS = "⚙️ Parámetros"

# Mapeo de la fila 10 (Bloque A) — columna de Excel (1-based) -> campo
BLOQUE_A = {
    1: "codigo", 2: "listing_nickname", 3: "ciudad", 4: "banco", 5: "modelo",
    6: "fecha_inicio", 7: "renta_base", 8: "comision_pct", 9: "iva_pct", 10: "irpf_pct",
    11: "limpieza_por_reserva", 12: "suministros_mes", 13: "comunidad_ibi_mes",
    14: "minut", 15: "akiles", 16: "amenities", 17: "pricelabs", 18: "guesty_fee",
    19: "extras", 20: "mobiliario_fin", 21: "propietario", 22: "nif", 23: "iban",
    24: "pasivo_base",
}
NUMERIC = {"renta_base", "comision_pct", "iva_pct", "irpf_pct", "limpieza_por_reserva",
           "suministros_mes", "comunidad_ibi_mes", "minut", "akiles", "amenities",
           "pricelabs", "guesty_fee", "extras", "mobiliario_fin", "pasivo_base"}


def sql_str(v):
    if v is None:
        return "NULL"
    return "'" + str(v).replace("'", "''") + "'"


def sql_num(v):
    if v is None or v == "":
        return "0"
    return repr(round(float(v), 4))


def norm_modelo(v):
    m = (v or "").strip().lower()
    return {"comisión": "comision", "comision": "comision",
            "subarriendo": "subarriendo", "titular": "titular"}.get(m, m)


def cell(ws, r, c):
    return ws.cell(r, c).value


def read_listings(ws):
    rows = []
    r = 11
    while cell(ws, r, 1):  # mientras haya código
        rec = {}
        for col, field in BLOQUE_A.items():
            v = cell(ws, r, col)
            if field == "modelo":
                v = norm_modelo(v)
            elif field == "fecha_inicio" and isinstance(v, dt.datetime):
                v = v.date().isoformat()
            rec[field] = v
        rows.append(rec)
        r += 1
    return rows


def read_general_expenses(ws):
    """Bloque B.1 — recurrentes: filas 20..31 (concepto, €/mes)."""
    rows = []
    r = 20
    while True:
        concepto = cell(ws, r, 1)
        importe = cell(ws, r, 2)
        if not concepto or str(concepto).upper().startswith("TOTAL"):
            break
        rows.append({"concepto": concepto, "importe_mes": importe})
        r += 1
    return rows


def read_events(ws):
    """Bloque C — eventos puntuales: desde fila 50 hasta que Año quede vacío."""
    rows = []
    r = 50
    while cell(ws, r, 1):  # Año
        anio = cell(ws, r, 1)
        mes = cell(ws, r, 2)
        rows.append({
            "anio": int(anio),
            "mes": int(mes),
            "propiedad_codigo": cell(ws, r, 3),
            "categoria": cell(ws, r, 4),
            "concepto": cell(ws, r, 5),
            "importe": cell(ws, r, 6),
            "notas": cell(ws, r, 7),
        })
        r += 1
    return rows


def main():
    xlsx = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_XLSX
    out = sys.argv[2] if len(sys.argv) > 2 else DEFAULT_OUT
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    ws = wb[PARAMS]

    listings = read_listings(ws)
    gexp = read_general_expenses(ws)
    events = read_events(ws)

    L = []
    L.append("-- seed.sql — generado por scripts/excel_to_seed.py (NO editar a mano)")
    L.append(f"-- Fuente: {xlsx} · hoja '{PARAMS}'")
    L.append("begin;")
    L.append("truncate table events, general_expenses, listings restart identity cascade;")
    L.append("")

    L.append("insert into listings (codigo, listing_nickname, ciudad, banco, modelo, fecha_inicio,")
    L.append("  renta_base, comision_pct, iva_pct, irpf_pct, limpieza_por_reserva, suministros_mes,")
    L.append("  comunidad_ibi_mes, minut, akiles, amenities, pricelabs, guesty_fee, extras,")
    L.append("  mobiliario_fin, propietario, nif, iban, pasivo_base) values")
    vals = []
    for x in listings:
        vals.append(
            "  (" + ", ".join([
                sql_str(x["codigo"]), sql_str(x["listing_nickname"]), sql_str(x["ciudad"]),
                sql_str(x["banco"]), sql_str(x["modelo"]), sql_str(x["fecha_inicio"]),
                sql_num(x["renta_base"]), sql_num(x["comision_pct"]), sql_num(x["iva_pct"]),
                sql_num(x["irpf_pct"]), sql_num(x["limpieza_por_reserva"]), sql_num(x["suministros_mes"]),
                sql_num(x["comunidad_ibi_mes"]), sql_num(x["minut"]), sql_num(x["akiles"]),
                sql_num(x["amenities"]), sql_num(x["pricelabs"]), sql_num(x["guesty_fee"]),
                sql_num(x["extras"]), sql_num(x["mobiliario_fin"]),
                # PII (propietario/nif/iban) NO se vuelca al repo → placeholders. Los reales viven en Supabase.
                "'PENDIENTE'", "'PENDIENTE'", "'PENDIENTE'", sql_num(x["pasivo_base"]),
            ]) + ")")
    L.append(",\n".join(vals) + ";")
    L.append("")

    L.append("insert into general_expenses (concepto, importe_mes) values")
    L.append(",\n".join(
        f"  ({sql_str(g['concepto'])}, {sql_num(g['importe_mes'])})" for g in gexp) + ";")
    L.append("")

    L.append("insert into events (anio, mes, propiedad_codigo, categoria, concepto, importe, notas) values")
    L.append(",\n".join(
        "  (" + ", ".join([
            str(e["anio"]), str(e["mes"]), sql_str(e["propiedad_codigo"]),
            sql_str(e["categoria"]), sql_str(e["concepto"]), sql_num(e["importe"]),
            sql_str(e["notas"]),
        ]) + ")" for e in events) + ";")
    L.append("")
    L.append("commit;")

    with open(out, "w", encoding="utf-8") as f:
        f.write("\n".join(L) + "\n")

    print(f"OK -> {out}")
    print(f"  listings: {len(listings)} | general_expenses: {len(gexp)} | events: {len(events)}")
    for x in listings:
        print(f"    {x['codigo']:9} modelo={x['modelo']:11} renta={x['renta_base']} comision%={x['comision_pct']}")


if __name__ == "__main__":
    main()
